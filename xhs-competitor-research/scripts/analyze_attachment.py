#!/usr/bin/env python3
"""Audit and profile a Xiaohongshu competitor-note attachment.

The script is read-only with respect to the source attachment. It writes derived
artifacts to an explicit output directory and never fetches network data. Its
statistics are intended to be cited by the xhs-competitor-research Skill.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import itertools
import json
import math
import re
import statistics
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple


SCHEMA_VERSION = 1
NULL_TEXT = {
    "",
    "-",
    "--",
    "null",
    "none",
    "nan",
    "n/a",
    "na",
    "不可见",
    "未知",
    "暂无",
    "无数据",
}

ALIASES: Dict[str, Tuple[str, ...]] = {
    "note_id": ("note_id", "笔记id", "笔记ID", "笔记编号", "小红书笔记id", "id"),
    "url": ("url", "canonical_url", "笔记链接", "链接", "笔记url", "笔记URL"),
    "title": ("title", "note_title", "标题", "笔记标题"),
    "body": (
        "body",
        "body_or_description",
        "正文",
        "笔记正文",
        "内容",
        "文案",
        "描述",
        "笔记内容",
    ),
    "tags": ("tags", "hashtags", "话题", "话题tag", "话题Tag", "话题标签", "标签"),
    "interaction_total": (
        "interaction_total",
        "engagement_total",
        "visible_subtotal",
        "总互动量",
        "互动量",
        "可见互动量",
    ),
    "likes": ("likes", "点赞", "点赞数", "点赞量", "interactions.likes"),
    "saves": (
        "saves",
        "collects",
        "收藏",
        "收藏数",
        "收藏量",
        "interactions.saves",
    ),
    "comments": ("comments", "评论", "评论数", "评论量", "interactions.comments"),
    "shares": ("shares", "分享", "分享数", "分享量", "interactions.shares"),
    "published_at": (
        "published_at",
        "published_at_normalized",
        "发布时间",
        "发布日期",
        "发布时刻",
    ),
    "author": ("author", "作者", "账号", "博主", "账号名称"),
    "account_type": ("account_type", "owner_type", "账号类型", "作者类型", "达人类型"),
    "follower_count": ("follower_count", "followers", "粉丝数", "账号粉丝数"),
    "content_type": ("content_type", "note_type", "笔记类型", "内容类型"),
    "image_count": ("image_count", "图片数量", "图片数", "visuals.image_count"),
    "image_type": (
        "image_type",
        "visual_type",
        "first_image_type",
        "图片类型",
        "视觉类型",
        "首图类型",
        "visuals.image_type",
    ),
    "opening_3s": ("opening_3s", "前三秒", "前3秒", "视频开头", "visuals.opening_3s"),
    "top_comments": ("top_comments", "评论内容", "头部评论", "热门评论"),
    "title_structure": ("title_structure", "标题结构", "标题公式"),
    "emotion_intensity": ("emotion_intensity", "情绪强度", "标题情绪强度"),
    "opening_hook": ("opening_hook", "开头钩子", "钩子类型"),
    "value_format": ("value_format", "价值呈现", "干货呈现方式"),
    "trust_device": ("trust_device", "信任手段", "信任建立手段"),
    "ending_cta": ("ending_cta", "结尾行动", "互动引导"),
    "functional_need": ("functional_need", "功能性需求"),
    "emotional_need": ("emotional_need", "情感性需求"),
    "social_need": ("social_need", "社交性需求"),
    "pain_point": ("pain_point", "痛点"),
    "pleasure_point": ("pleasure_point", "爽点"),
    "itch_point": ("itch_point", "痒点"),
}

NUMERIC_FIELDS = {
    "interaction_total",
    "likes",
    "saves",
    "comments",
    "shares",
    "follower_count",
    "image_count",
    "emotion_intensity",
}

TEXT_FIELDS = set(ALIASES) - NUMERIC_FIELDS - {"tags"}
INTERACTION_FIELDS = ("interaction_total", "likes", "saves", "comments", "shares")
SEMANTIC_FIELDS = (
    "content_type",
    "title_structure",
    "emotion_intensity",
    "opening_hook",
    "value_format",
    "trust_device",
    "ending_cta",
    "functional_need",
    "emotional_need",
    "social_need",
    "pain_point",
    "pleasure_point",
    "itch_point",
)

STOPWORDS = {
    "一个",
    "一些",
    "这个",
    "那个",
    "我们",
    "你们",
    "他们",
    "自己",
    "就是",
    "还是",
    "没有",
    "可以",
    "真的",
    "非常",
    "什么",
    "怎么",
    "为什么",
    "因为",
    "所以",
    "如果",
    "但是",
    "而且",
    "以及",
    "已经",
    "今天",
    "现在",
    "感觉",
    "分享",
    "小红书",
}

EMOTION_WORDS = {
    "焦虑",
    "害怕",
    "崩溃",
    "后悔",
    "震惊",
    "惊喜",
    "绝了",
    "救命",
    "必看",
    "爱了",
    "喜欢",
    "治愈",
    "幸福",
    "愤怒",
    "踩雷",
    "避雷",
    "血亏",
    "真香",
}
CONTRAST_WORDS = {"对比", "反差", "居然", "竟然", "反而", "但是", "却", "vs", "VS", "前后"}
IDENTITY_WORDS = {"女生", "男生", "宝妈", "学生", "新手", "打工人", "职场人", "上班族", "00后", "90后"}
URGENCY_WORDS = {"最后", "仅限", "赶紧", "马上", "别再", "现在", "一定要", "必看", "错过"}
PAIN_WORDS = {"痛点", "困扰", "焦虑", "踩坑", "踩雷", "失败", "崩溃", "后悔", "难受"}
BENEFIT_WORDS = {"省钱", "省时", "变好", "提升", "解决", "学会", "轻松", "效率", "效果", "方法"}


def normalize_header(value: str) -> str:
    return re.sub(r"[\s_\-—（）()【】\[\]./\\]+", "", value.strip().lower())


ALIAS_LOOKUP: Dict[str, str] = {}
for canonical_name, alias_values in ALIASES.items():
    for alias_value in alias_values:
        ALIAS_LOOKUP[normalize_header(alias_value)] = canonical_name


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, (list, tuple, dict)):
        return len(value) == 0
    return str(value).strip().lower() in NULL_TEXT


def parse_number(value: Any) -> Tuple[Optional[float], bool]:
    if is_blank(value):
        return None, False
    if isinstance(value, bool):
        return float(value), False
    if isinstance(value, (int, float)):
        number = float(value)
        return (number, False) if math.isfinite(number) else (None, False)

    text = str(value).strip()
    approximate = bool(re.search(r"[+＋~～约]", text))
    text = re.sub(r"[+＋~～约]", "", text)
    text = re.sub(r"[￥¥$,，\s]", "", text)
    text = re.sub(r"(?:次|个|条|赞|收藏|评论|分享)$", "", text)
    multiplier = 1.0
    suffixes = (("亿", 100_000_000.0), ("万", 10_000.0), ("w", 10_000.0), ("千", 1_000.0), ("k", 1_000.0))
    lowered = text.lower()
    for suffix, factor in suffixes:
        if lowered.endswith(suffix):
            multiplier = factor
            text = text[: -len(suffix)]
            break
    try:
        number = float(text) * multiplier
    except ValueError:
        return None, approximate
    if not math.isfinite(number) or number < 0:
        return None, approximate
    return number, approximate


def parse_tags(value: Any) -> List[str]:
    if is_blank(value):
        return []
    if isinstance(value, (list, tuple, set)):
        raw_tags = [str(item) for item in value if not is_blank(item)]
    else:
        text = str(value).strip()
        hashtag_matches = re.findall(r"#\s*([^#，,；;|、\n\r]+)", text)
        if hashtag_matches:
            raw_tags = hashtag_matches
        else:
            raw_tags = re.split(r"[，,；;|、\n\r]+", text)

    result: List[str] = []
    seen = set()
    for raw_tag in raw_tags:
        tag = re.sub(r"^#+", "", raw_tag.strip()).strip()
        if not tag:
            continue
        normalized = tag.lower()
        if normalized in seen:
            continue
        seen.add(normalized)
        result.append(tag)
    return result


def flatten_record(record: Dict[str, Any]) -> Dict[str, Any]:
    result: Dict[str, Any] = {}

    def visit(prefix: str, value: Any) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                child_prefix = f"{prefix}.{key}" if prefix else str(key)
                visit(child_prefix, child)
            return
        result[prefix] = value
        leaf = prefix.rsplit(".", 1)[-1]
        result.setdefault(leaf, value)

    for raw_key, raw_value in record.items():
        visit(str(raw_key), raw_value)
    return result


def decode_csv(path: Path) -> Tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 编码无法识别；请另存为 UTF-8 CSV")


def load_records(path_text: str, input_format: Optional[str], sheet_name: Optional[str]) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    if path_text == "-":
        fmt = (input_format or "csv").lower()
        raw_text = sys.stdin.read()
        source_name = "stdin"
        source_bytes = raw_text.encode("utf-8")
        encoding = "utf-8"
    else:
        source_path = Path(path_text).expanduser().resolve()
        if not source_path.is_file():
            raise ValueError(f"附件不存在或不是文件：{source_path}")
        fmt = (input_format or source_path.suffix.lstrip(".")).lower()
        source_name = source_path.name
        source_bytes = source_path.read_bytes()
        raw_text = ""
        encoding = None

    meta: Dict[str, Any] = {
        "file_name": source_name,
        "format": fmt,
        "sheet": sheet_name,
        "sha256": hashlib.sha256(source_bytes).hexdigest(),
    }

    if fmt == "csv":
        if path_text != "-":
            raw_text, encoding = decode_csv(Path(path_text).expanduser().resolve())
        reader = csv.DictReader(io.StringIO(raw_text))
        if not reader.fieldnames:
            raise ValueError("CSV 缺少表头")
        records = [dict(row) for row in reader]
        meta["encoding"] = encoding
        return records, meta

    if fmt == "json":
        if path_text != "-":
            raw_text = source_bytes.decode("utf-8-sig")
        payload = json.loads(raw_text)
        if isinstance(payload, dict):
            for key in ("records", "data", "notes"):
                if isinstance(payload.get(key), list):
                    payload = payload[key]
                    break
        if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
            raise ValueError("JSON 必须是对象数组，或包含 records/data/notes 对象数组")
        return list(payload), meta

    if fmt in {"jsonl", "ndjson"}:
        if path_text != "-":
            raw_text = source_bytes.decode("utf-8-sig")
        records = []
        for line_number, line in enumerate(raw_text.splitlines(), start=1):
            if not line.strip():
                continue
            item = json.loads(line)
            if not isinstance(item, dict):
                raise ValueError(f"JSONL 第 {line_number} 行不是对象")
            records.append(item)
        return records, meta

    if fmt == "xlsx":
        if path_text == "-":
            raise ValueError("XLSX 不支持从 stdin 读取")
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:
            raise ValueError("读取 XLSX 需要 openpyxl；请安装依赖或只读另存为 UTF-8 CSV") from exc
        workbook = load_workbook(Path(path_text).expanduser().resolve(), read_only=True, data_only=True)
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"工作表不存在：{sheet_name}；可用：{', '.join(workbook.sheetnames)}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = ["" if value is None else str(value) for value in next(rows)]
        except StopIteration as exc:
            raise ValueError("XLSX 工作表为空") from exc
        if not any(headers):
            raise ValueError("XLSX 首行缺少表头")
        records = [dict(zip(headers, row)) for row in rows if any(not is_blank(value) for value in row)]
        meta["sheet"] = sheet.title
        meta["available_sheets"] = workbook.sheetnames
        return records, meta

    raise ValueError("仅支持 CSV、JSON、JSONL 和 XLSX")


def canonicalize_record(raw_record: Dict[str, Any], source_row: int) -> Dict[str, Any]:
    flattened = flatten_record(raw_record)
    canonical: Dict[str, Any] = {}
    field_sources: Dict[str, str] = {}
    warnings: List[str] = []
    approximate_fields: List[str] = []

    for raw_key, value in flattened.items():
        canonical_name = ALIAS_LOOKUP.get(normalize_header(str(raw_key)))
        if canonical_name is None or is_blank(value):
            continue
        if canonical_name in canonical and not is_blank(canonical[canonical_name]):
            if str(canonical[canonical_name]).strip() != str(value).strip():
                warnings.append(
                    f"字段别名冲突：{field_sources[canonical_name]} 与 {raw_key} 均映射为 {canonical_name}"
                )
            continue
        canonical[canonical_name] = value
        field_sources[canonical_name] = str(raw_key)

    result: Dict[str, Any] = {"source_row": source_row, "warnings": warnings}
    for field in TEXT_FIELDS:
        value = canonical.get(field)
        if isinstance(value, (list, dict)):
            result[field] = json.dumps(value, ensure_ascii=False)
        else:
            result[field] = None if is_blank(value) else str(value).strip()
    result["tags"] = parse_tags(canonical.get("tags"))

    for field in NUMERIC_FIELDS:
        value = canonical.get(field)
        parsed, approximate = parse_number(value)
        result[field] = parsed
        if approximate:
            approximate_fields.append(field)
        if not is_blank(value) and parsed is None:
            warnings.append(f"无法解析数值字段 {field}={value!r}")
    result["approximate_fields"] = approximate_fields
    result["field_sources"] = field_sources
    return result


def normalize_url(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    return value.strip().split("#", 1)[0]


def note_key(record: Dict[str, Any]) -> str:
    if record.get("note_id"):
        return f"id:{record['note_id']}"
    normalized_url = normalize_url(record.get("url"))
    if normalized_url:
        return f"url:{normalized_url}"
    title = record.get("title") or ""
    body = record.get("body") or ""
    if title or body:
        digest = hashlib.sha1(f"{title}\x1f{body}".encode("utf-8")).hexdigest()
        return f"text:{digest}"
    return f"row:{record['source_row']}"


def completeness(record: Dict[str, Any]) -> int:
    score = 0
    for field in ALIASES:
        value = record.get(field)
        if not is_blank(value):
            score += 1
    return score


def deduplicate(records: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    grouped: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for record in records:
        record["note_key"] = note_key(record)
        grouped[record["note_key"]].append(record)

    unique_records: List[Dict[str, Any]] = []
    duplicate_rows = 0
    conflict_groups: List[Dict[str, Any]] = []
    critical_fields = {"title", "body", "interaction_total", "likes", "saves", "comments", "shares"}

    for key in sorted(grouped, key=lambda item: min(row["source_row"] for row in grouped[item])):
        group = grouped[key]
        duplicate_rows += max(0, len(group) - 1)
        base = dict(max(group, key=lambda row: (completeness(row), -row["source_row"])))
        base["warnings"] = []
        base["approximate_fields"] = []
        base["source_rows"] = sorted(row["source_row"] for row in group)
        merged_tags: List[str] = []
        seen_tags = set()
        approximate_fields = set()
        conflicts = set()

        for row in group:
            for tag in row.get("tags", []):
                normalized = tag.lower()
                if normalized not in seen_tags:
                    seen_tags.add(normalized)
                    merged_tags.append(tag)
            approximate_fields.update(row.get("approximate_fields", []))
            for field in ALIASES:
                if field == "tags":
                    continue
                current = base.get(field)
                candidate = row.get(field)
                if is_blank(current) and not is_blank(candidate):
                    base[field] = candidate
                elif not is_blank(current) and not is_blank(candidate) and str(current).strip() != str(candidate).strip():
                    conflicts.add(field)
            base.setdefault("warnings", []).extend(row.get("warnings", []))

        base["tags"] = merged_tags
        base["approximate_fields"] = sorted(approximate_fields)
        base["conflict_fields"] = sorted(conflicts)
        base["exclude_from_metric"] = bool(conflicts & critical_fields)
        if conflicts:
            conflict_groups.append(
                {"note_key": key, "source_rows": base["source_rows"], "fields": sorted(conflicts)}
            )
        unique_records.append(base)

    return unique_records, {
        "duplicate_rows": duplicate_rows,
        "duplicate_groups": sum(1 for group in grouped.values() if len(group) > 1),
        "conflict_groups": conflict_groups,
    }


def nonblank_count(records: Sequence[Dict[str, Any]], field: str) -> int:
    return sum(1 for record in records if not is_blank(record.get(field)))


def choose_interaction_metric(
    records: List[Dict[str, Any]], explicit_fields: Optional[str], coverage_threshold: float
) -> Dict[str, Any]:
    total = len(records)
    coverage = {
        field: (nonblank_count(records, field) / total if total else 0.0)
        for field in INTERACTION_FIELDS
    }
    warnings: List[str] = []

    if explicit_fields and explicit_fields.lower() != "auto":
        selected: List[str] = []
        for raw_field in explicit_fields.split(","):
            normalized = normalize_header(raw_field)
            canonical = ALIAS_LOOKUP.get(normalized, raw_field.strip())
            if canonical not in INTERACTION_FIELDS:
                raise ValueError(f"不支持的互动字段：{raw_field}")
            if canonical not in selected:
                selected.append(canonical)
        if "interaction_total" in selected and len(selected) > 1:
            raise ValueError("总互动字段不能与点赞、收藏、评论或分享重复相加")
        mode = "explicit"
    elif coverage["interaction_total"] >= coverage_threshold:
        selected = ["interaction_total"]
        mode = "reported_total"
    else:
        component_candidates = [
            field for field in ("likes", "saves", "comments", "shares") if coverage[field] >= coverage_threshold
        ]
        if component_candidates:
            selected = component_candidates
            mode = "component_sum"
        else:
            ranked = sorted(
                (field for field in ("likes", "saves", "comments", "shares") if coverage[field] > 0),
                key=lambda field: (-coverage[field], field),
            )
            if ranked and coverage[ranked[0]] >= 0.5:
                selected = [ranked[0]]
                mode = "single_component_fallback"
                warnings.append("没有互动字段达到统一覆盖阈值，使用覆盖率最高的单项指标；结论需降级")
            else:
                selected = []
                mode = "unavailable"

    valid_count = 0
    approximate_count = 0
    for record in records:
        record["interaction_score"] = None
        if record.get("exclude_from_metric") or not selected:
            continue
        values = [record.get(field) for field in selected]
        if any(value is None for value in values):
            continue
        score = values[0] if selected == ["interaction_total"] else sum(values)
        record["interaction_score"] = float(score)
        record["interaction_metric"] = "+".join(selected)
        valid_count += 1
        if any(field in record.get("approximate_fields", []) for field in selected):
            approximate_count += 1

    return {
        "status": "ready" if selected else "unavailable",
        "mode": mode,
        "selected_fields": selected,
        "definition": " + ".join(selected) if selected else None,
        "field_coverage": {field: round(value, 4) for field, value in coverage.items()},
        "complete_rows": valid_count,
        "complete_rate": round(valid_count / total, 4) if total else 0.0,
        "approximate_rows": approximate_count,
        "warnings": warnings,
    }


def assign_tiers(records: List[Dict[str, Any]], minimum_rows: int) -> Dict[str, Any]:
    eligible = [record for record in records if record.get("interaction_score") is not None]
    if len(eligible) < minimum_rows:
        for record in records:
            record["tier"] = "unclassified"
        return {
            "enabled": False,
            "reason": f"完整互动样本 {len(eligible)} 条，少于最低要求 {minimum_rows} 条",
            "eligible_rows": len(eligible),
            "counts": {},
            "boundary_ties": [],
        }
    if len({record["interaction_score"] for record in eligible}) < 2:
        for record in records:
            record["tier"] = "unclassified"
        return {
            "enabled": False,
            "reason": "互动指标没有差异，无法形成有效分层",
            "eligible_rows": len(eligible),
            "counts": {},
            "boundary_ties": [],
        }

    ordered = sorted(eligible, key=lambda record: (-record["interaction_score"], record["note_key"]))
    head_count = int(math.ceil(len(ordered) * 0.10))
    waist_count = int(math.ceil(len(ordered) * 0.30))
    boundary_ties: List[Dict[str, Any]] = []

    for index, record in enumerate(ordered):
        if index < head_count:
            record["tier"] = "head"
        elif index < head_count + waist_count:
            record["tier"] = "waist"
        else:
            record["tier"] = "tail"

    for record in records:
        record.setdefault("tier", "unclassified")

    boundaries = ((head_count, "head/waist"), (head_count + waist_count, "waist/tail"))
    for boundary_index, boundary_name in boundaries:
        if 0 < boundary_index < len(ordered):
            left_value = ordered[boundary_index - 1]["interaction_score"]
            right_value = ordered[boundary_index]["interaction_score"]
            if left_value == right_value:
                boundary_ties.append({"boundary": boundary_name, "score": left_value})

    counts = Counter(record["tier"] for record in ordered)
    return {
        "enabled": True,
        "reason": None,
        "eligible_rows": len(ordered),
        "counts": {tier: counts.get(tier, 0) for tier in ("head", "waist", "tail")},
        "cutoffs": {"head": head_count, "waist": waist_count, "tail": len(ordered) - head_count - waist_count},
        "boundary_ties": boundary_ties,
        "warning": "边界并列值按稳定 note_key 切分，层级差异仅作探索" if boundary_ties else None,
    }


def compact_chars(value: Optional[str]) -> str:
    return re.sub(r"\s+", "", value or "")


def count_hits(text: str, words: Iterable[str]) -> int:
    lowered = text.lower()
    return sum(lowered.count(word.lower()) for word in words)


def parse_published_at(value: Optional[str]) -> Tuple[Optional[int], Optional[str], Optional[str]]:
    if not value:
        return None, None, None
    parsed: Optional[datetime] = None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        text = str(value).strip()
        normalized = text.replace("年", "-").replace("月", "-").replace("日", " ").replace("/", "-")
        normalized = re.sub(r"\s+", " ", normalized).strip()
        try:
            parsed = datetime.fromisoformat(normalized.replace("Z", "+00:00"))
        except ValueError:
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
                try:
                    parsed = datetime.strptime(normalized, fmt)
                    break
                except ValueError:
                    continue
    if parsed is None:
        return None, None, None
    weekday_names = ("周一", "周二", "周三", "周四", "周五", "周六", "周日")
    hour = parsed.hour
    if 0 <= hour < 6:
        period = "凌晨"
    elif hour < 12:
        period = "上午"
    elif hour < 18:
        period = "下午"
    else:
        period = "晚间"
    return hour, weekday_names[parsed.weekday()], period


def derive_features(record: Dict[str, Any]) -> None:
    title = record.get("title") or ""
    body = record.get("body") or ""
    title_compact = compact_chars(title)
    body_compact = compact_chars(body)
    paragraphs = [part.strip() for part in re.split(r"\r?\n+", body) if part.strip()]
    punctuation_count = title.count("!") + title.count("！") + title.count("?") + title.count("？")
    emotion_hits = count_hits(title, EMOTION_WORDS)
    emotion_types = sum(
        1
        for word_set in (EMOTION_WORDS, PAIN_WORDS, BENEFIT_WORDS, URGENCY_WORDS)
        if count_hits(title, word_set) > 0
    )
    emotion_score = 1
    if emotion_hits >= 1:
        emotion_score += 1
    if emotion_types >= 2:
        emotion_score += 1
    if punctuation_count >= 2:
        emotion_score += 1
    if emotion_hits >= 4:
        emotion_score += 1

    opening_50 = body_compact[:50]
    ending_80 = body_compact[-80:]
    opening_signals = {
        "question": bool(re.search(r"[?？]|为什么|怎么|是否", opening_50)),
        "number": bool(re.search(r"\d", opening_50)),
        "conflict": count_hits(opening_50, CONTRAST_WORDS) > 0,
        "pain": count_hits(opening_50, PAIN_WORDS) > 0,
        "benefit": count_hits(opening_50, BENEFIT_WORDS) > 0,
    }
    cta_signals = {
        "comment": "评论" in ending_80 or "留言" in ending_80,
        "save": "收藏" in ending_80,
        "share": "分享" in ending_80 or "转发" in ending_80,
        "follow": "关注" in ending_80,
        "like": "点赞" in ending_80,
    }

    record.update(
        {
            "title_chars": len(title_compact),
            "body_chars": len(body_compact),
            "paragraph_count": len(paragraphs),
            "tag_count": len(record.get("tags", [])),
            "unique_tag_count": len({tag.lower() for tag in record.get("tags", [])}),
            "title_has_number": int(bool(re.search(r"\d|[一二三四五六七八九十百千万两]", title))),
            "title_question_count": title.count("?") + title.count("？"),
            "title_exclamation_count": title.count("!") + title.count("！"),
            "title_bracket_count": sum(title.count(char) for char in "【】[]（）()"),
            "title_contrast_signal": int(count_hits(title, CONTRAST_WORDS) > 0),
            "title_identity_signal": int(count_hits(title, IDENTITY_WORDS) > 0),
            "title_urgency_signal": int(count_hits(title, URGENCY_WORDS) > 0),
            "emotion_intensity_heuristic": min(emotion_score, 5),
            "opening_50": opening_50,
            "opening_hook_count": sum(int(value) for value in opening_signals.values()),
            "opening_signals": [key for key, value in opening_signals.items() if value],
            "ending_80": ending_80,
            "cta_count": sum(int(value) for value in cta_signals.values()),
            "cta_signals": [key for key, value in cta_signals.items() if value],
        }
    )
    hour, weekday, period = parse_published_at(record.get("published_at"))
    record["published_hour"] = hour
    record["published_weekday"] = weekday
    record["published_period"] = period


def average_ranks(values: Sequence[float]) -> List[float]:
    indexed = sorted(enumerate(values), key=lambda pair: pair[1])
    ranks = [0.0] * len(values)
    index = 0
    while index < len(indexed):
        end = index + 1
        while end < len(indexed) and indexed[end][1] == indexed[index][1]:
            end += 1
        average = (index + 1 + end) / 2.0
        for position in range(index, end):
            ranks[indexed[position][0]] = average
        index = end
    return ranks


def pearson(x_values: Sequence[float], y_values: Sequence[float]) -> Optional[float]:
    if len(x_values) != len(y_values) or len(x_values) < 2:
        return None
    x_mean = statistics.fmean(x_values)
    y_mean = statistics.fmean(y_values)
    numerator = sum((x - x_mean) * (y - y_mean) for x, y in zip(x_values, y_values))
    x_denominator = sum((x - x_mean) ** 2 for x in x_values)
    y_denominator = sum((y - y_mean) ** 2 for y in y_values)
    denominator = math.sqrt(x_denominator * y_denominator)
    return numerator / denominator if denominator else None


def spearman(x_values: Sequence[float], y_values: Sequence[float]) -> Optional[float]:
    return pearson(average_ranks(x_values), average_ranks(y_values))


def quantile(values: Sequence[float], probability: float) -> Optional[float]:
    if not values:
        return None
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    position = (len(ordered) - 1) * probability
    lower = int(math.floor(position))
    upper = int(math.ceil(position))
    if lower == upper:
        return ordered[lower]
    fraction = position - lower
    return ordered[lower] * (1 - fraction) + ordered[upper] * fraction


NUMERIC_FEATURES = (
    "title_chars",
    "body_chars",
    "paragraph_count",
    "tag_count",
    "unique_tag_count",
    "image_count",
    "follower_count",
    "title_has_number",
    "title_question_count",
    "title_exclamation_count",
    "title_bracket_count",
    "title_contrast_signal",
    "title_identity_signal",
    "title_urgency_signal",
    "emotion_intensity",
    "emotion_intensity_heuristic",
    "opening_hook_count",
    "cta_count",
)

CATEGORICAL_FEATURES = (
    "account_type",
    "content_type",
    "image_type",
    "published_weekday",
    "published_period",
    "title_structure",
    "opening_hook",
    "value_format",
    "trust_device",
    "ending_cta",
    "functional_need",
    "emotional_need",
    "social_need",
)


def numeric_associations(records: Sequence[Dict[str, Any]], minimum_n: int) -> Dict[str, Any]:
    metric_rows = [record for record in records if record.get("interaction_score") is not None]
    results = []
    excluded = []
    for feature in NUMERIC_FEATURES:
        pairs = [
            (float(record[feature]), float(record["interaction_score"]))
            for record in metric_rows
            if record.get(feature) is not None
        ]
        missing_rate = 1 - len(pairs) / len(metric_rows) if metric_rows else 1.0
        if len(pairs) < minimum_n:
            excluded.append({"feature": feature, "reason": f"成对样本 {len(pairs)} < {minimum_n}"})
            continue
        if missing_rate > 0.30:
            excluded.append({"feature": feature, "reason": f"缺失率 {missing_rate:.1%} > 30%"})
            continue
        x_values = [pair[0] for pair in pairs]
        y_values = [pair[1] for pair in pairs]
        if len(set(x_values)) < 2 or len(set(y_values)) < 2:
            excluded.append({"feature": feature, "reason": "变量没有足够变化"})
            continue
        coefficient = spearman(x_values, y_values)
        if coefficient is None:
            continue
        results.append(
            {
                "feature": feature,
                "n": len(pairs),
                "rho": round(coefficient, 4),
                "abs_rho": round(abs(coefficient), 4),
                "missing_rate": round(missing_rate, 4),
                "wording": "与互动表现正相关"
                if coefficient > 0
                else "与互动表现负相关"
                if coefficient < 0
                else "未发现单调相关",
            }
        )
    results.sort(key=lambda item: (-item["abs_rho"], item["feature"]))
    return {
        "numeric": results,
        "top_variables": [item for item in results if item["abs_rho"] >= 0.10][:5],
        "excluded": excluded,
    }


def categorical_associations(records: Sequence[Dict[str, Any]], minimum_group_n: int) -> List[Dict[str, Any]]:
    metric_rows = [record for record in records if record.get("interaction_score") is not None]
    all_scores = [float(record["interaction_score"]) for record in metric_rows]
    overall_median = statistics.median(all_scores) if all_scores else None
    results = []
    for feature in CATEGORICAL_FEATURES:
        grouped: Dict[str, List[float]] = defaultdict(list)
        for record in metric_rows:
            value = record.get(feature)
            if is_blank(value):
                continue
            grouped[str(value).strip()].append(float(record["interaction_score"]))
        eligible_groups = {key: values for key, values in grouped.items() if len(values) >= minimum_group_n}
        if len(eligible_groups) < 2 or len(eligible_groups) > 20:
            continue
        group_rows = []
        for key, values in sorted(eligible_groups.items()):
            median_value = statistics.median(values)
            lift = median_value / overall_median if overall_median not in (None, 0) else None
            group_rows.append(
                {
                    "group": key,
                    "n": len(values),
                    "median": round(median_value, 4),
                    "p25": round(quantile(values, 0.25) or 0, 4),
                    "p75": round(quantile(values, 0.75) or 0, 4),
                    "median_lift_vs_overall": round(lift, 4) if lift is not None else None,
                }
            )
        results.append(
            {
                "feature": feature,
                "overall_median": round(overall_median, 4) if overall_median is not None else None,
                "groups": group_rows,
                "omitted_small_groups": sorted(key for key, values in grouped.items() if len(values) < minimum_group_n),
            }
        )
    return results


def summarize_values(values: Sequence[float]) -> Dict[str, Any]:
    if not values:
        return {"n": 0, "mean": None, "median": None, "min": None, "max": None}
    return {
        "n": len(values),
        "mean": round(statistics.fmean(values), 4),
        "median": round(statistics.median(values), 4),
        "min": round(min(values), 4),
        "max": round(max(values), 4),
    }


def tier_numeric_summary(records: Sequence[Dict[str, Any]], tiering_enabled: bool) -> Dict[str, Any]:
    if not tiering_enabled:
        return {}
    result: Dict[str, Any] = {}
    for feature in NUMERIC_FEATURES:
        feature_summary = {}
        for tier in ("head", "waist", "tail"):
            values = [float(record[feature]) for record in records if record.get("tier") == tier and record.get(feature) is not None]
            feature_summary[tier] = summarize_values(values)
        if any(item["n"] for item in feature_summary.values()):
            result[feature] = feature_summary
    return result


def load_jieba() -> Any:
    try:
        import jieba  # type: ignore

        jieba.setLogLevel(30)
        return jieba
    except ImportError:
        return None


def tokenize_documents(text: str, jieba_module: Any) -> set:
    tokens = set()
    if jieba_module is not None:
        candidates = jieba_module.lcut(text)
        for candidate in candidates:
            token = candidate.strip().lower()
            if 2 <= len(token) <= 20 and token not in STOPWORDS and re.search(r"[a-z0-9\u4e00-\u9fff]", token):
                tokens.add(token)
        return tokens

    for latin in re.findall(r"[A-Za-z][A-Za-z0-9+.-]{1,19}", text):
        tokens.add(latin.lower())
    for segment in re.findall(r"[\u4e00-\u9fff]{2,}", text):
        for size in (2, 3, 4):
            if len(segment) < size:
                continue
            for index in range(len(segment) - size + 1):
                token = segment[index : index + size]
                if token not in STOPWORDS:
                    tokens.add(token)
    return tokens


def frequency_rows(counter: Counter, document_count: int, examples: Dict[str, List[str]], limit: int) -> List[Dict[str, Any]]:
    rows = []
    for token, count in sorted(counter.items(), key=lambda item: (-item[1], item[0]))[:limit]:
        rows.append(
            {
                "token": token,
                "documents": count,
                "rate": round(count / document_count, 4) if document_count else 0.0,
                "example_note_keys": examples.get(token, [])[:3],
            }
        )
    return rows


def lexical_analysis(records: Sequence[Dict[str, Any]], top_tags: int, tiering_enabled: bool) -> Dict[str, Any]:
    jieba_module = load_jieba()
    method = "jieba_document_frequency" if jieba_module is not None else "chinese_character_ngrams_2_4_document_frequency"
    warning = None if jieba_module is not None else "未安装 jieba；候选词使用中文 2–4 字符 n-gram，报告中不得直接称为人工语义关键词"
    minimum_head_lift = 1.05

    title_counter: Counter = Counter()
    body_counter: Counter = Counter()
    title_examples: Dict[str, List[str]] = defaultdict(list)
    body_examples: Dict[str, List[str]] = defaultdict(list)
    title_documents = 0
    body_documents = 0
    token_sets_by_key: Dict[str, set] = {}

    for record in records:
        title = record.get("title") or ""
        body = record.get("body") or ""
        title_tokens = tokenize_documents(title, jieba_module) if title else set()
        body_tokens = tokenize_documents(body, jieba_module) if body else set()
        token_sets_by_key[record["note_key"]] = title_tokens | body_tokens
        if title:
            title_documents += 1
            title_counter.update(title_tokens)
            for token in title_tokens:
                title_examples[token].append(record["note_key"])
        if body:
            body_documents += 1
            body_counter.update(body_tokens)
            for token in body_tokens:
                body_examples[token].append(record["note_key"])

    tag_counter: Counter = Counter()
    tag_examples: Dict[str, List[str]] = defaultdict(list)
    pair_counter: Counter = Counter()
    for record in records:
        normalized_tags = sorted({tag.strip().lower() for tag in record.get("tags", []) if tag.strip()})
        tag_counter.update(normalized_tags)
        for tag in normalized_tags:
            tag_examples[tag].append(record["note_key"])
        pair_counter.update(itertools.combinations(normalized_tags, 2))

    head_specific_terms = []
    distinctive_tags = []
    if tiering_enabled:
        head_records = [record for record in records if record.get("tier") == "head"]
        nonhead_records = [record for record in records if record.get("tier") in {"waist", "tail"}]
        head_n = len(head_records)
        nonhead_n = len(nonhead_records)
        all_tokens = set().union(*(token_sets_by_key[record["note_key"]] for record in records)) if records else set()
        for token in all_tokens:
            head_count = sum(token in token_sets_by_key[record["note_key"]] for record in head_records)
            nonhead_count = sum(token in token_sets_by_key[record["note_key"]] for record in nonhead_records)
            if head_count < 2:
                continue
            head_rate = head_count / head_n if head_n else 0
            nonhead_rate = nonhead_count / nonhead_n if nonhead_n else 0
            smoothed_lift = ((head_count + 0.5) / (head_n + 1)) / ((nonhead_count + 0.5) / (nonhead_n + 1))
            if head_rate <= nonhead_rate or smoothed_lift < minimum_head_lift:
                continue
            head_specific_terms.append(
                {
                    "token": token,
                    "head_count": head_count,
                    "head_denominator": head_n,
                    "head_rate": round(head_rate, 4),
                    "nonhead_count": nonhead_count,
                    "nonhead_denominator": nonhead_n,
                    "nonhead_rate": round(nonhead_rate, 4),
                    "smoothed_lift": round(smoothed_lift, 4),
                }
            )
        head_specific_terms.sort(key=lambda item: (-item["smoothed_lift"], -item["head_count"], item["token"]))

        all_tags = set(tag_counter)
        for tag in all_tags:
            head_count = sum(tag in {item.lower() for item in record.get("tags", [])} for record in head_records)
            nonhead_count = sum(tag in {item.lower() for item in record.get("tags", [])} for record in nonhead_records)
            if head_count < 2:
                continue
            head_rate = head_count / head_n if head_n else 0
            nonhead_rate = nonhead_count / nonhead_n if nonhead_n else 0
            smoothed_lift = ((head_count + 0.5) / (head_n + 1)) / ((nonhead_count + 0.5) / (nonhead_n + 1))
            if head_rate <= nonhead_rate or smoothed_lift < minimum_head_lift:
                continue
            distinctive_tags.append(
                {
                    "tag": tag,
                    "head_count": head_count,
                    "head_denominator": head_n,
                    "head_rate": round(head_rate, 4),
                    "nonhead_count": nonhead_count,
                    "nonhead_denominator": nonhead_n,
                    "nonhead_rate": round(nonhead_rate, 4),
                    "smoothed_lift": round(smoothed_lift, 4),
                }
            )
        distinctive_tags.sort(key=lambda item: (-item["smoothed_lift"], -item["head_count"], item["tag"]))

    tag_rows = []
    for tag, count in sorted(tag_counter.items(), key=lambda item: (-item[1], item[0]))[:top_tags]:
        tag_rows.append(
            {
                "tag": tag,
                "documents": count,
                "rate": round(count / len(records), 4) if records else 0.0,
                "example_note_keys": tag_examples[tag][:3],
            }
        )

    pair_rows = [
        {"tags": list(pair), "documents": count}
        for pair, count in sorted(pair_counter.items(), key=lambda item: (-item[1], item[0]))[:30]
    ]
    return {
        "method": method,
        "warning": warning,
        "token_interpretation": (
            "document-frequency terms"
            if jieba_module is not None
            else "character n-gram candidates requiring manual semantic review"
        ),
        "valid_text_documents": {
            "title": title_documents,
            "body": body_documents,
            "title_or_body": sum(
                bool(record.get("title") or record.get("body")) for record in records
            ),
            "denominator": len(records),
        },
        "stopword_rule": "built-in list; tokens shorter than 2 characters are excluded",
        "stopwords": sorted(STOPWORDS),
        "head_enrichment_rule": (
            f"head_count >= 2, head_rate > nonhead_rate, smoothed_lift >= {minimum_head_lift}"
        ),
        "title_top30": frequency_rows(title_counter, title_documents, title_examples, 30),
        "body_top50": frequency_rows(body_counter, body_documents, body_examples, 50),
        "tags": tag_rows,
        "tag_pairs": pair_rows,
        "head_specific_terms": head_specific_terms[:30],
        "head_distinctive_tags": distinctive_tags[:20],
    }


def field_audit(records: Sequence[Dict[str, Any]]) -> Dict[str, Any]:
    total = len(records)
    coverage = {}
    for field in ALIASES:
        count = nonblank_count(records, field)
        coverage[field] = {
            "nonblank": count,
            "denominator": total,
            "rate": round(count / total, 4) if total else 0.0,
        }
    return coverage


def capability_status(coverage: Dict[str, Dict[str, Any]], fields: Sequence[str], threshold: float = 0.70) -> str:
    rates = [coverage[field]["rate"] for field in fields]
    if rates and all(rate >= threshold for rate in rates):
        return "ready"
    if any(rate > 0 for rate in rates):
        return "partial"
    return "unavailable"


def build_capabilities(
    coverage: Dict[str, Dict[str, Any]], tiering: Dict[str, Any], associations: Dict[str, Any]
) -> Dict[str, Any]:
    return {
        "title_body_patterns": {
            "status": "ready" if coverage["title"]["rate"] > 0 or coverage["body"]["rate"] > 0 else "unavailable",
            "required": ["title or body"],
        },
        "tag_analysis": {"status": capability_status(coverage, ["tags"]), "required": ["tags"]},
        "interaction_tiering": {
            "status": "ready" if tiering["enabled"] else "unavailable",
            "required": ["consistent interaction metric", "at least 30 complete rows", "metric variance"],
            "reason": tiering.get("reason"),
        },
        "numeric_associations": {
            "status": "ready" if associations["numeric"] else "unavailable",
            "required": ["interaction metric", "numeric feature", "at least 30 paired rows"],
        },
        "publish_time": {"status": capability_status(coverage, ["published_at"]), "required": ["published_at"]},
        "account_type": {"status": capability_status(coverage, ["account_type"]), "required": ["account_type"]},
        "visual": {
            "status": "ready"
            if capability_status(coverage, ["image_count"]) == "ready" or capability_status(coverage, ["image_type"]) == "ready"
            else "partial"
            if coverage["image_count"]["rate"] > 0 or coverage["image_type"]["rate"] > 0
            else "unavailable",
            "required": ["image_count and/or image_type"],
        },
        "first_three_seconds": {"status": capability_status(coverage, ["opening_3s"]), "required": ["opening_3s or time-coded transcript"]},
        "comment_drivers": {"status": capability_status(coverage, ["top_comments"]), "required": ["top_comments"]},
        "semantic_needs": {
            "status": "ready"
            if any(coverage[field]["rate"] >= 0.70 for field in ("functional_need", "emotional_need", "social_need"))
            else "requires_coding",
            "required": ["saved semantic coding"],
        },
        "platform_low_competition_tags": {
            "status": "unavailable",
            "required": ["sample-external platform supply or competition baseline"],
        },
        "expected_interaction_volume": {
            "status": "unavailable",
            "required": ["comparable historical validation data"],
        },
    }


def cleaned_row(record: Dict[str, Any]) -> Dict[str, Any]:
    fields = (
        "note_key",
        "source_rows",
        "note_id",
        "url",
        "title",
        "body",
        "tags",
        "published_at",
        "author",
        "account_type",
        "follower_count",
        "content_type",
        "image_count",
        "image_type",
        "opening_3s",
        "top_comments",
        "interaction_total",
        "likes",
        "saves",
        "comments",
        "shares",
        "interaction_metric",
        "interaction_score",
        "tier",
        "title_chars",
        "body_chars",
        "paragraph_count",
        "tag_count",
        "unique_tag_count",
        "title_has_number",
        "title_question_count",
        "title_exclamation_count",
        "title_bracket_count",
        "title_contrast_signal",
        "title_identity_signal",
        "title_urgency_signal",
        "emotion_intensity_heuristic",
        "opening_50",
        "opening_signals",
        "ending_80",
        "cta_signals",
        "published_hour",
        "published_weekday",
        "published_period",
        "conflict_fields",
        "approximate_fields",
        "warnings",
    )
    result = {}
    for field in fields:
        value = record.get(field)
        if isinstance(value, list):
            if field == "tags":
                result[field] = " ".join(f"#{item}" for item in value)
            else:
                result[field] = json.dumps(value, ensure_ascii=False)
        else:
            result[field] = value
    return result


def semantic_template_row(record: Dict[str, Any]) -> Dict[str, Any]:
    result = {
        "note_key": record["note_key"],
        "source_rows": json.dumps(record.get("source_rows", []), ensure_ascii=False),
        "note_id": record.get("note_id"),
        "url": record.get("url"),
        "title": record.get("title"),
        "body": record.get("body"),
        "tags": " ".join(f"#{item}" for item in record.get("tags", [])),
        "interaction_total": record.get("interaction_total"),
        "likes": record.get("likes"),
        "saves": record.get("saves"),
        "comments": record.get("comments"),
        "shares": record.get("shares"),
        "tier": record.get("tier"),
    }
    for field in SEMANTIC_FIELDS:
        result[field] = record.get(field) or ""
    result["evidence_excerpt"] = ""
    result["coding_notes"] = ""
    return result


def write_csv(path: Path, rows: List[Dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"无法写出空表：{path.name}")
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def analyze(args: argparse.Namespace) -> Dict[str, Any]:
    raw_records, source_meta = load_records(args.input, args.input_format, args.sheet)
    if not raw_records:
        raise ValueError("附件没有数据行")

    source_index_start = 2 if source_meta["format"] in {"csv", "xlsx"} else 1
    canonical_records = [
        canonicalize_record(record, index)
        for index, record in enumerate(raw_records, start=source_index_start)
    ]
    records, dedupe_meta = deduplicate(canonical_records)
    for record in records:
        derive_features(record)

    coverage = field_audit(records)
    interaction_metric = choose_interaction_metric(records, args.interaction_fields, args.coverage_threshold)
    tiering = assign_tiers(records, args.minimum_tier_rows)
    associations = numeric_associations(records, args.minimum_correlation_rows)
    associations["categorical"] = categorical_associations(records, args.minimum_group_rows)
    tiering["numeric_summaries"] = tier_numeric_summary(records, tiering["enabled"])
    lexical = lexical_analysis(records, args.top_tags, tiering["enabled"])
    capabilities = build_capabilities(coverage, tiering, associations)

    if len(records) < 30:
        strength = "descriptive_only"
    elif len(records) < 100:
        strength = "exploratory"
    elif len(records) <= 300:
        strength = "fuller_exploratory"
    else:
        strength = "large_sample_exploratory"

    invalid_text_rows = [
        {"note_key": record["note_key"], "source_rows": record.get("source_rows", [])}
        for record in records
        if not record.get("title") and not record.get("body")
    ]
    warnings = list(interaction_metric["warnings"])
    if lexical.get("warning"):
        warnings.append(lexical["warning"])
    if tiering.get("warning"):
        warnings.append(tiering["warning"])
    if invalid_text_rows:
        warnings.append(f"{len(invalid_text_rows)} 条记录同时缺少标题和正文")
    semantic_prefilled_rows = sum(
        any(not is_blank(record.get(field)) for field in SEMANTIC_FIELDS)
        for record in records
    )
    if semantic_prefilled_rows < len(records):
        warnings.append(
            "脚本只完成可复算的定量审计；semantic-coding-template.csv 仍需逐篇编码、抽查并建立证据账本"
        )

    summary = {
        "schema_version": SCHEMA_VERSION,
        "source": source_meta,
        "meta": {
            "input_rows": len(raw_records),
            "unique_notes": len(records),
            "analysis_strength": strength,
            "script_scope": "quantitative_audit_only",
            "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        },
        "data_quality": {
            "field_coverage": coverage,
            "duplicates": dedupe_meta,
            "invalid_text_rows": invalid_text_rows,
            "record_warning_count": sum(len(record.get("warnings", [])) for record in records),
            "warnings": warnings,
            "capabilities": capabilities,
        },
        "interaction_metric": interaction_metric,
        "tiering": tiering,
        "associations": associations,
        "lexical": lexical,
        "semantic_coding": {
            "template_rows": len(records),
            "prefilled_rows": semantic_prefilled_rows,
            "status": (
                "prefilled_requires_audit"
                if semantic_prefilled_rows == len(records)
                else "requires_coding_and_audit"
            ),
            "evidence_ledger_generated": False,
        },
        "evidence_contract": {
            "percentage_fields": ["numerator", "denominator", "sample_scope"],
            "case_key": "note_key with source_rows",
            "causal_language_allowed": False,
        },
    }

    output_dir = Path(args.output_dir).expanduser().resolve()
    summary_path = output_dir / "analysis-summary.json"
    cleaned_path = output_dir / "cleaned-notes.csv"
    coding_path = output_dir / "semantic-coding-template.csv"
    existing_outputs = [
        path for path in (summary_path, cleaned_path, coding_path) if path.exists()
    ]
    if existing_outputs and not args.overwrite:
        existing_names = ", ".join(path.name for path in existing_outputs)
        raise ValueError(
            f"派生输出已存在：{existing_names}；请使用新的 --output-dir，"
            "或在明确允许覆盖时增加 --overwrite"
        )
    output_dir.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_csv(cleaned_path, [cleaned_row(record) for record in records])
    write_csv(coding_path, [semantic_template_row(record) for record in records])

    return {
        "status": "ok",
        "source_file": source_meta["file_name"],
        "input_rows": len(raw_records),
        "unique_notes": len(records),
        "interaction_metric": interaction_metric["definition"],
        "tiering_enabled": tiering["enabled"],
        "analysis_strength": strength,
        "script_scope": "quantitative_audit_only",
        "semantic_coding_status": summary["semantic_coding"]["status"],
        "summary_path": str(summary_path),
        "cleaned_notes_path": str(cleaned_path),
        "semantic_coding_template_path": str(coding_path),
        "warnings": warnings,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="审计并分析小红书竞品笔记附件")
    parser.add_argument("input", help="CSV/JSON/JSONL/XLSX 路径；使用 - 从 stdin 读取")
    parser.add_argument("--output-dir", required=True, help="派生结果输出目录，不覆盖源附件")
    parser.add_argument("--input-format", choices=("csv", "json", "jsonl", "ndjson", "xlsx"))
    parser.add_argument("--sheet", help="XLSX 工作表名称；省略时使用活动工作表")
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="明确允许覆盖输出目录中已有的三个派生文件",
    )
    parser.add_argument(
        "--interaction-fields",
        default="auto",
        help="auto，或逗号分隔的总互动/点赞/收藏/评论/分享字段",
    )
    parser.add_argument("--coverage-threshold", type=float, default=0.90, help="自动选择互动字段的最低覆盖率")
    parser.add_argument("--minimum-tier-rows", type=int, default=30)
    parser.add_argument("--minimum-correlation-rows", type=int, default=30)
    parser.add_argument("--minimum-group-rows", type=int, default=5)
    parser.add_argument("--top-tags", type=int, default=50)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    if not 0 < args.coverage_threshold <= 1:
        parser.error("--coverage-threshold 必须在 0 到 1 之间")
    for field in ("minimum_tier_rows", "minimum_correlation_rows", "minimum_group_rows", "top_tags"):
        if getattr(args, field) <= 0:
            parser.error(f"--{field.replace('_', '-')} 必须为正整数")
    try:
        result = analyze(args)
    except (ValueError, OSError, json.JSONDecodeError, csv.Error) as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
